"""第七阶段汇总的冻结边界、数字来源与隐私检查。"""

from __future__ import annotations

import json
from pathlib import Path

import nbformat
import pytest
from openpyxl import load_workbook

from src.reporting.stage7 import build_summary

ROOT = Path(__file__).resolve().parents[1]


def test_missing_frozen_artifact_reports_chinese_error(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="缺少已验收产物"):
        build_summary(tmp_path, include_strategy=False)


def test_summary_reuses_frozen_metrics_and_training_strategy() -> None:
    summary = build_summary(ROOT)
    source = json.loads((ROOT / "reports/stage5_summary.json").read_text(encoding="utf-8"))
    assert summary["primary_model"] == "Logistic 评分卡"
    assert summary["best_tree_challenger"] == "XGBoost"
    assert summary["frozen_manifest_sha256"] == source["frozen_manifest_sha256"]
    for split in ("validation_metrics", "test_metrics_frozen_audit"):
        for actual, expected in zip(summary[split], source[split], strict=True):
            assert actual["model"] == expected["model"]
            assert actual["auc"] == expected["auc"]
            assert actual["ks"] == expected["ks"]
    assert summary["strategy_at_600_train_only"]["cutoff"] == 600
    assert 0 <= summary["strategy_at_600_train_only"]["pass_rate"] <= 1


def test_comprehensive_workbook_contains_aggregates_only() -> None:
    path = ROOT / "reports/excel/stage7_comprehensive_report.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert len(workbook.sheetnames) == 13
    assert workbook.sheetnames[0] == "结论总览"
    assert {"分箱WOE", "评分卡明细", "验证集模型比较", "冻结测试审计"}.issubset(workbook.sheetnames)
    for sheet in workbook:
        headers = [str(cell.value) for cell in sheet[5] if cell.value is not None]
        assert "source_row_id" not in headers
    overview = list(workbook["结论总览"].values)
    assert len([row for row in overview if len(row) > 1 and row[1] == "验证集"]) == 3
    assert len([row for row in overview if len(row) > 1 and row[1] == "测试集审计"]) == 3
    workbook.close()


def test_eighth_notebook_keeps_aggregate_outputs_without_personal_case_output() -> None:
    notebook = nbformat.read(ROOT / "notebooks/08_完整案例与风险策略.ipynb", as_version=4)
    code = [cell for cell in notebook.cells if cell.cell_type == "code"]
    assert len(code) == 3
    assert code[0].execution_count is not None
    assert code[1].execution_count is None
    assert code[1].outputs == []
    assert code[2].execution_count is not None
    assert all(output.output_type != "error" for cell in code for output in cell.outputs)
    assert "strategy_at_cutoff" in "\n".join(cell.source for cell in code)


def test_public_stage5_artifacts_exclude_individual_rows() -> None:
    workbook = load_workbook(ROOT / "reports/excel/stage5_model_comparison_public.xlsx", read_only=True, data_only=True)
    assert len(workbook.sheetnames) == 15
    assert "SHAP个案" not in workbook.sheetnames
    assert "解释客户" not in workbook.sheetnames
    private_columns = {"source_row_id", "row_position", "raw_feature_value", "model_input_value"}
    for sheet in workbook:
        headers = {str(cell.value) for cell in sheet[5] if cell.value is not None}
        assert not headers & private_columns
    workbook.close()

    notebook = nbformat.read(ROOT / "notebooks/07_SHAP模型解释.ipynb", as_version=4)
    code = [cell for cell in notebook.cells if cell.cell_type == "code"]
    assert code[1].execution_count is None
    assert code[1].outputs == []
    assert "source_row_id" not in "\n".join(str(output) for cell in code for output in cell.outputs)
