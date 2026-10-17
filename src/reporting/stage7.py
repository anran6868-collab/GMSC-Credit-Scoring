"""从已验收产物生成可追溯的第七阶段汇总，不训练或评价模型。"""

from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook

from src.dashboard.service import load_manifest, strategy_at_cutoff, strategy_population

ROOT = Path(__file__).resolve().parents[2]
METRICS = ("auc", "ks", "pr_auc", "f1", "brier", "log_loss")
MODELS = ("Logistic 评分卡", "Random Forest", "XGBoost")


def _required(path: Path) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"缺少第七阶段所需的已验收产物：{path}")
    return path


def _read_json(path: Path) -> dict:
    return json.loads(_required(path).read_text(encoding="utf-8"))


def _clean_json(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, dict):
        return {key: _clean_json(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_clean_json(item) for item in value]
    return value


def _source_table(path: Path, sheet: str, header_row: int = 1) -> list[list]:
    source = load_workbook(_required(path), read_only=True, data_only=True)
    try:
        if sheet not in source.sheetnames:
            raise ValueError(f"工作簿缺少工作表：{path.name} / {sheet}")
        rows = [list(row) for row in source[sheet].iter_rows(min_row=header_row, values_only=True)]
        rows = [row for row in rows if any(value is not None for value in row)]
        if not rows or len(rows) > 1000:
            raise ValueError(f"工作表为空或超出汇总行数上限：{sheet}")
        if any("source_row_id" in str(value) for value in rows[0]):
            raise ValueError(f"工作表含逐行客户标识，不得进入综合报告：{sheet}")
        return rows
    finally:
        source.close()


def _verify_metrics(expected: list[dict], rows: list[list], label: str) -> None:
    header = [str(value) for value in rows[0]]
    model_col = header.index("model")
    actual = {str(row[model_col]): row for row in rows[1:] if row[model_col] is not None}
    for item in expected:
        name = item["model"]
        if name not in actual:
            raise ValueError(f"{label}缺少模型：{name}")
        for metric in METRICS:
            number = float(item[metric])
            workbook_number = float(actual[name][header.index(metric)])
            if not math.isfinite(number) or not math.isclose(number, workbook_number, abs_tol=1e-10):
                raise ValueError(f"{label}中{name}的{metric}与阶段报告不一致")


def build_summary(root: Path = ROOT, *, include_strategy: bool = True) -> dict:
    """读取冻结清单与阶段报告；只在训练集计算已验收的策略口径。"""
    root = Path(root)
    frozen = load_manifest(root)
    frozen_hash = _read_json(root / "models/stage5_frozen_manifest.json")["sha256"]
    stage3 = _read_json(root / "reports/stage3_summary.json")
    stage4 = _read_json(root / "reports/stage4_summary.json")
    stage5 = _read_json(root / "reports/stage5_summary.json")
    if stage5["frozen_manifest_sha256"] != frozen_hash:
        raise ValueError("第五阶段汇总与冻结清单的 SHA-256 不一致")
    if stage5["primary_model"] != "Logistic 评分卡" or stage5["test_used_for_selection"] is not False:
        raise ValueError("第五阶段模型路线或测试集使用口径不符合项目要求")
    if stage4["selected_feature_count"] != len(stage4["selected_features"]):
        raise ValueError("第四阶段入模变量数量不一致")
    sizes = stage3["split_sizes"]
    if set(sizes) != {"train", "validation", "test"} or any(not isinstance(value, int) or value <= 0 for value in sizes.values()):
        raise ValueError("分层划分规模不完整")
    total = sum(sizes.values())
    for split, expected_share in (("train", 0.70), ("validation", 0.15), ("test", 0.15)):
        if not math.isclose(sizes[split] / total, expected_share, abs_tol=1 / total):
            raise ValueError("分层划分比例与项目要求不一致")
    quality_sizes = _source_table(root / "reports/excel/data_quality.xlsx", "划分规模")
    if len(quality_sizes) < 2 or dict(zip(quality_sizes[0], quality_sizes[1])) != sizes:
        raise ValueError("第二阶段数据质量工作簿与划分规模不一致")
    comparisons = root / "reports/excel/stage5_model_comparison.xlsx"
    validation = _source_table(comparisons, "验证集比较", 5)
    test = _source_table(comparisons, "测试集审计", 5)
    _verify_metrics(stage5["validation_metrics"], validation, "验证集比较")
    _verify_metrics(stage5["test_metrics_frozen_audit"], test, "测试集审计")
    if {item["model"] for item in stage5["validation_metrics"]} != set(MODELS):
        raise ValueError("验证集模型集合不完整")
    strategy = None
    if include_strategy:
        strategy = strategy_at_cutoff(strategy_population(root), 600)
    return {
        "project": "GMSC 金融风控评分卡建模项目",
        "target": "未来两年严重逾期概率",
        "split_sizes": stage3["split_sizes"],
        "candidate_feature_count": stage3["candidate_feature_count"],
        "selected_features": stage4["selected_features"],
        "selected_scheme": stage4["selected_scheme"],
        "score_parameters": {"base_score": 600, "odds_good_bad": 20, "pdo": 20},
        "grade_boundaries": stage4["grade_boundaries"],
        "primary_model": stage5["primary_model"],
        "best_tree_challenger": stage5["best_tree_challenger"],
        "validation_metrics": _clean_json(stage5["validation_metrics"]),
        "test_metrics_frozen_audit": _clean_json(stage5["test_metrics_frozen_audit"]),
        "shap_additivity_max_error": stage5["shap_additivity_max_error"],
        "frozen_manifest_sha256": frozen_hash,
        "strategy_at_600_train_only": strategy,
        "sources": [
            "reports/stage3_summary.json",
            "reports/stage4_summary.json",
            "reports/stage5_summary.json",
            "models/stage5_frozen_manifest.json",
            "reports/excel/stage5_model_comparison.xlsx",
        ],
    }


SHEETS = (
    ("data_quality.xlsx", "缺失与分布", 1, "数据质量"),
    ("data_quality.xlsx", "异常与特殊值", 1, "异常规则"),
    ("stage3_woe_iv_analysis.xlsx", "特征定义", 1, "特征定义"),
    ("stage3_woe_iv_analysis.xlsx", "训练分箱明细", 1, "分箱WOE"),
    ("stage3_woe_iv_analysis.xlsx", "IV_PSI与初筛", 1, "IV与稳定性"),
    ("stage4_logistic_scorecard.xlsx", "变量筛选轨迹", 1, "特征筛选"),
    ("stage4_logistic_scorecard.xlsx", "评分卡规则", 1, "评分卡明细"),
    ("stage4_logistic_scorecard.xlsx", "等级统计", 1, "风险等级"),
    ("stage5_model_comparison.xlsx", "验证集比较", 5, "验证集模型比较"),
    ("stage5_model_comparison.xlsx", "测试集审计", 5, "冻结测试审计"),
    ("stage5_model_comparison.xlsx", "SHAP全局", 5, "SHAP汇总"),
)


def write_workbook_data(summary: dict, root: Path = ROOT) -> Path:
    """导出经过白名单筛选的汇总表，供本地制表工具排版。"""
    root = Path(root)
    tables = [{
        "sheet": "结论总览",
        "columns": ["模型", "数据集", "AUC", "KS", "PR-AUC", "F1", "Brier", "Log Loss"],
        "rows": [
            [item["model"], split, *(item[key] for key in METRICS)]
            for split, items in (("验证集", summary["validation_metrics"]), ("测试集审计", summary["test_metrics_frozen_audit"]))
            for item in items
        ],
    }]
    for source_file, source_sheet, header_row, target_sheet in SHEETS:
        rows = _source_table(root / "reports/excel" / source_file, source_sheet, header_row)
        columns = [str(value) if value is not None else f"字段{i + 1}" for i, value in enumerate(rows[0])]
        records = [
            [None if isinstance(value, float) and not math.isfinite(value) else value for value in row]
            for row in rows[1:]
        ]
        tables.append({"sheet": target_sheet, "columns": columns, "rows": records, "source": f"{source_file} / {source_sheet}"})
    tables.append({
        "sheet": "来源与边界",
        "columns": ["项目", "内容"],
        "rows": [
            ["冻结清单 SHA-256", summary["frozen_manifest_sha256"]],
            ["训练集规模", summary["split_sizes"]["train"]],
            ["验证集规模", summary["split_sizes"]["validation"]],
            ["测试集规模", summary["split_sizes"]["test"]],
            ["数值核对", "验证集与冻结测试审计同第五阶段工作簿一致"],
            ["应用边界", "教学与研究；不代表真实授信审批"],
            ["策略边界", "训练集样本内模拟；结果可能偏乐观"],
        ],
    })
    output = root / "reports/stage7_workbook_tables.json"
    output.write_text(json.dumps(tables, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    return output


def run(root: Path = ROOT) -> dict:
    summary = build_summary(root)
    output = Path(root) / "reports/stage7_summary.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    write_workbook_data(summary, root)
    return summary


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2, allow_nan=False))
